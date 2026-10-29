"""openpyxl Excel report builders."""
import re
from io import BytesIO

from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .builders import Cell

NAVY = "3B1466"  # theme purple
STATUS_FILL = {
    "COMPLIED": "D1E7DD",
    "EXPIRING": "D1E7DD",
    "INCOMPLETE": "FFF3CD",
    "NOT_COMPLIED": "F8D7DA",
    "NA": "E9ECEF",
    "PERCENT": "EFE7FA",
}
THIN = Side(style="thin", color="ADB5BD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _sheet_title(name, used):
    title = re.sub(r"[\[\]\*\?/\\:]", "", name)[:31] or "Sheet"
    base, n = title, 2
    while title in used:
        title = f"{base[:28]}_{n}"
        n += 1
    used.add(title)
    return title


def _write_header(ws, report, heading=None):
    ws.append([settings.SCHOOL_NAME])
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws.append([f"FRAMES — {settings.SYSTEM_FULL_NAME}"])
    ws["A2"].font = Font(italic=True, size=9, color="495057")
    ws.append([report.title + (f" — {heading}" if heading else "")])
    ws["A3"].font = Font(bold=True, size=12, color=NAVY)
    ws.append(["Filters: " + ("; ".join(f"{k}: {v}" for k, v in report.filters) or "None")])
    ws.append([f"Generated: {report.generated_at:%Y-%m-%d %H:%M} by {report.generated_by}"])
    ws.append([])
    return ws.max_row + 1


def _value(v):
    text = v.text if isinstance(v, Cell) else v
    if isinstance(text, str) and text.endswith("%"):
        try:
            return float(text.rstrip("%")) / 100
        except ValueError:
            return text
    return text


def render_xlsx(report):
    wb = Workbook()
    used = set()

    ws = wb.active
    ws.title = _sheet_title("Summary", used)
    _write_header(ws, report)
    for label, value in report.summary:
        ws.append([label, value])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 60

    for table in report.tables:
        ws = wb.create_sheet(_sheet_title(table.sheet_name or table.heading, used))
        start = _write_header(ws, report, table.heading)
        ws.append(table.headers)
        for col in range(1, len(table.headers) + 1):
            cell = ws.cell(start, col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER
        for row in table.rows:
            ws.append([_value(v) for v in row])
            r = ws.max_row
            for c, v in enumerate(row, start=1):
                cell = ws.cell(r, c)
                cell.border = BORDER
                if isinstance(v, Cell):
                    if v.status in STATUS_FILL:
                        cell.fill = PatternFill("solid", fgColor=STATUS_FILL[v.status])
                    if v.bold:
                        cell.font = Font(bold=True)
                    if isinstance(cell.value, float) and v.text.endswith("%"):
                        cell.number_format = "0.00%"
                    if table.font_size < 7 and c > 2:
                        cell.alignment = Alignment(horizontal="center")
        if not table.rows:
            ws.append(["No records match the selected filters."])
        # column widths
        widths = table.col_widths or [1] * len(table.headers)
        for i, w in enumerate(widths, start=1):
            letter = get_column_letter(i)
            if table.font_size < 7:
                ws.column_dimensions[letter].width = 30 if i == 1 else (8 if i == 2 else 7)
            else:
                ws.column_dimensions[letter].width = max(8, min(70, w * 9))
        if table.freeze:
            col = re.match(r"[A-Z]+", table.freeze).group(0)
            ws.freeze_panes = f"{col}{start + 1}"
        else:
            ws.freeze_panes = f"A{start + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
